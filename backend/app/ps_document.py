"""
Генерация ПС в DOCX/HTML по образцу заполненного макета и приказу 446н.
"""
from __future__ import annotations

import io
import html
import re
from typing import Any

from .db.raw_models import StandardRaw

REF_LABELS = (
    ("ОКЗ", "okz_units"),
    ("ЕКС", "etks_units"),
    ("ОКПДТР", "okpdtr_units"),
    ("ОКСО", "okso_units"),
)

# A4 21 см; поля слева/сверху/снизу 2 см, справа 1 см → контент 18 см
CONTENT_W = 18.0


def _codes(value) -> list[str]:
    if not value:
        return []
    if isinstance(value, list):
        return [str(v) for v in value if v]
    return [str(value)]


def _units(value) -> list[dict]:
    if not value:
        return []
    result = []
    for item in value:
        if isinstance(item, dict):
            result.append({"code": item.get("code") or "", "name": item.get("name") or ""})
        else:
            result.append({"code": str(item), "name": ""})
    return result


def _job_titles_text(titles: list[str]) -> str:
    cleaned = [t.strip() for t in (titles or []) if t and t.strip()]
    if not cleaned:
        return ""
    return " / или / ".join(cleaned)


def _job_titles_map(titles: list[str]) -> str:
    """В функциональной карте — через запятую (как в образце)."""
    cleaned = [t.strip() for t in (titles or []) if t and t.strip()]
    return ", ".join(cleaned)


def _split_okz_pairs(code: str, name: str) -> list[tuple[str, str]]:
    codes = [c.strip() for c in re.split(r"[;,]", code or "") if c.strip()]
    names = [n.strip() for n in re.split(r"[;]", name or "") if n.strip()]
    if not codes and not names:
        return [("", "")]
    if len(codes) <= 1 and len(names) <= 1:
        return [(code or "", name or "")]
    pairs = []
    for i in range(max(len(codes), len(names))):
        pairs.append((codes[i] if i < len(codes) else "", names[i] if i < len(names) else ""))
    return pairs or [("", "")]


def _format_order_stamp(approval_date: str, order_number: str) -> list[str]:
    date = (approval_date or "").strip() or "«___» _______ ____"
    number = (order_number or "").strip() or "___"
    return [
        "УТВЕРЖДЕН",
        "приказом Министерства",
        "труда и социальной защиты",
        "Российской Федерации",
        f"от {date} г. № {number}",
    ]


def build_ps_document_dict(std: StandardRaw) -> dict[str, Any]:
    okved = _units(getattr(std, "okved_units", None))
    if not okved:
        okved = [{"code": c, "name": ""} for c in _codes(std.okved_codes)]

    gfs = []
    for gf in std.generalized_functions:
        pfs = []
        for pf in gf.particular_functions:
            pfs.append(
                {
                    "code": pf.code or "",
                    "name": pf.name or "",
                    "sub_qualification": pf.sub_qualification or "",
                    "labor_actions": [la.text for la in (pf.labor_actions or []) if la.text],
                    "required_skills": [s.text for s in (pf.skills or []) if s.text],
                    "necessary_knowledges": [k.text for k in (pf.knowledges or []) if k.text],
                    "other_characteristics": getattr(pf, "other_characteristics", None) or "",
                }
            )
        gfs.append(
            {
                "code": gf.code or "",
                "name": gf.name or "",
                "level": gf.level or "",
                "possible_job_titles": gf.possible_job_titles or [],
                "education_training": getattr(gf, "education_training", None) or "",
                "practical_experience": getattr(gf, "practical_experience", None) or "",
                "special_admission": getattr(gf, "special_admission", None) or "",
                "other_characteristics": getattr(gf, "other_characteristics", None) or "",
                "okz_units": _units(getattr(gf, "okz_units", None))
                or [{"code": c, "name": ""} for c in _codes(gf.okz_codes)],
                "okpdtr_units": _units(getattr(gf, "okpdtr_units", None))
                or [{"code": c, "name": ""} for c in _codes(gf.okpdtr_codes)],
                "okso_units": _units(getattr(gf, "okso_units", None))
                or [{"code": c, "name": ""} for c in _codes(gf.okso_codes)],
                "etks_units": _units(getattr(gf, "etks_units", None)),
                "particular_functions": pfs,
            }
        )

    return {
        "reg_number": std.reg_number or "",
        "name": std.name or "",
        "order_number": std.order_number or "",
        "approval_date": std.approval_date or "",
        "ps_code": getattr(std, "ps_code", None) or "",
        "kind_activity": std.kind_activity or "",
        "purpose": std.purpose or "",
        "opd_code": getattr(std, "opd_code", None) or std.professional_area_code or "",
        "opd_name": getattr(std, "opd_name", None) or "",
        "okz_group_code": getattr(std, "okz_group_code", None) or "",
        "okz_group_name": getattr(std, "okz_group_name", None) or "",
        "okved_units": okved,
        "developer_org": getattr(std, "developer_org", None) or "",
        "developer_head": getattr(std, "developer_head", None) or "",
        "co_developers": getattr(std, "co_developers", None) or [],
        "abbreviations": getattr(std, "abbreviations", None) or [],
        "generalized_functions": gfs,
    }


def _build_map_rows(doc: dict[str, Any]) -> list[dict]:
    rows: list[dict] = []
    for gf in doc["generalized_functions"]:
        titles = _job_titles_map(gf["possible_job_titles"])
        pfs = gf["particular_functions"] or [{"name": "", "code": "", "sub_qualification": ""}]
        for pf in pfs:
            rows.append(
                {
                    "otf_code": gf["code"] or "",
                    "otf_name": gf["name"] or "",
                    "otf_level": gf["level"] or "",
                    "titles": titles,
                    "tf_name": pf["name"] or "",
                    "tf_code": pf["code"] or "",
                    "tf_level": pf["sub_qualification"] or "",
                }
            )
    return rows


def render_ps_html(doc: dict[str, Any]) -> str:
    def esc(v) -> str:
        return html.escape(str(v or ""))

    def rows_per_item(label: str, items: list[str]) -> str:
        if not items:
            return f"<tr><td class='lab'>{esc(label)}</td><td></td></tr>"
        return "".join(
            f"<tr><td class='lab'>{esc(label)}</td><td>{esc(x)}</td></tr>" for x in items
        )

    stamp_lines = _format_order_stamp(doc["approval_date"], doc["order_number"])
    stamp_html = "<br/>".join(esc(x) for x in stamp_lines)

    okz_pairs = _split_okz_pairs(doc["okz_group_code"], doc["okz_group_name"])
    okz_rows_html = []
    for i in range(0, len(okz_pairs), 2):
        left = okz_pairs[i]
        right = okz_pairs[i + 1] if i + 1 < len(okz_pairs) else ("", "")
        okz_rows_html.append(
            f"<tr><td>{esc(left[0])}</td><td>{esc(left[1])}</td>"
            f"<td>{esc(right[0])}</td><td>{esc(right[1])}</td></tr>"
        )
    okz_rows_html.append(
        "<tr class='cap'><td>(код ОКЗ)</td><td>(наименование)</td>"
        "<td>(код ОКЗ)</td><td>(наименование)</td></tr>"
    )

    okved_body = "".join(
        f"<tr><td>{esc(u['code'])}</td><td>{esc(u['name'])}</td></tr>" for u in doc["okved_units"]
    ) or "<tr><td></td><td></td></tr>"
    okved_body += (
        "<tr class='cap'><td>(код ОКВЭД)</td>"
        "<td>(наименование вида экономической деятельности)</td></tr>"
    )

    # Карта с rowspan по ОТФ
    map_rows = _build_map_rows(doc)
    map_html_parts = []
    i = 0
    while i < len(map_rows):
        j = i + 1
        while (
            j < len(map_rows)
            and map_rows[j]["otf_code"] == map_rows[i]["otf_code"]
            and map_rows[j]["otf_name"] == map_rows[i]["otf_name"]
        ):
            j += 1
        span = j - i
        for k in range(i, j):
            row = map_rows[k]
            if k == i:
                map_html_parts.append(
                    "<tr>"
                    f"<td rowspan='{span}'>{esc(row['otf_code'])}</td>"
                    f"<td rowspan='{span}'>{esc(row['otf_name'])}</td>"
                    f"<td rowspan='{span}'>{esc(row['otf_level'])}</td>"
                    f"<td rowspan='{span}'>{esc(row['titles'])}</td>"
                    f"<td>{esc(row['tf_name'])}</td>"
                    f"<td>{esc(row['tf_code'])}</td>"
                    f"<td>{esc(row['tf_level'])}</td></tr>"
                )
            else:
                map_html_parts.append(
                    "<tr>"
                    f"<td>{esc(row['tf_name'])}</td>"
                    f"<td>{esc(row['tf_code'])}</td>"
                    f"<td>{esc(row['tf_level'])}</td></tr>"
                )
        i = j
    map_html = "".join(map_html_parts) or "<tr>" + ("<td></td>" * 7) + "</tr>"

    toc = [
        "I. Общие сведения",
        "II. Описание трудовых функций, входящих в профессиональный стандарт "
        "(функциональная карта вида профессиональной деятельности)",
        "III. Характеристика обобщенных трудовых функций",
    ]
    for idx, gf in enumerate(doc["generalized_functions"], start=1):
        toc.append(f"3.{idx}. Обобщенная трудовая функция «{gf['name']}»")
    toc += [
        "IV. Сведения об организациях – разработчиках профессионального стандарта",
        "4.1. Ответственная организация-разработчик",
        "4.2. Наименования организаций-разработчиков",
        "V. Сокращения, используемые в профессиональном стандарте",
    ]
    toc_html = "".join(f"<div>{esc(x)}</div>" for x in toc)

    gf_html = []
    for gi, gf in enumerate(doc["generalized_functions"], start=1):
        ref_rows = []
        for label, key in REF_LABELS:
            units = gf.get(key) or []
            if not units:
                ref_rows.append(f"<tr><td>{esc(label)}</td><td></td><td></td></tr>")
            else:
                for u in units:
                    ref_rows.append(
                        f"<tr><td>{esc(label)}</td><td>{esc(u.get('code'))}</td>"
                        f"<td>{esc(u.get('name'))}</td></tr>"
                    )
        pf_parts = []
        for j, pf in enumerate(gf["particular_functions"], start=1):
            body = (
                rows_per_item("Трудовые действия", pf["labor_actions"])
                + rows_per_item("Необходимые умения", pf["required_skills"])
                + rows_per_item("Необходимые знания", pf["necessary_knowledges"])
                + "<tr><td class='lab'>Другие характеристики</td>"
                f"<td>{esc(pf['other_characteristics'] or '-')}</td></tr>"
            )
            pf_parts.append(
                f"""
                <h3>3.{gi}.{j}. Трудовая функция</h3>
                <table class="meta6">
                  <tr>
                    <td class="lab">Наименование</td><td>{esc(pf['name'])}</td>
                    <td class="lab">Код</td><td>{esc(pf['code'])}</td>
                    <td class="lab">Уровень (подуровень) квалификации</td>
                    <td>{esc(pf['sub_qualification'])}</td>
                  </tr>
                </table>
                <table>{body}</table>
                """
            )
        gf_html.append(
            f"""
            <h2>3.{gi}. Обобщенная трудовая функция «{esc(gf['name'])}»</h2>
            <table class="meta6">
              <tr>
                <td class="lab">Наименование</td><td>{esc(gf['name'])}</td>
                <td class="lab">Код</td><td>{esc(gf['code'])}</td>
                <td class="lab">Уровень квалификации</td><td>{esc(gf['level'])}</td>
              </tr>
            </table>
            <table>
              <tr><td class="lab">Возможные наименования должностей, профессий рабочих</td>
                  <td>{esc(_job_titles_text(gf['possible_job_titles']))}</td></tr>
            </table>
            <p class="sub">Пути достижения квалификации</p>
            <table>
              <tr><td class="lab">Образование и обучение</td>
                  <td>{esc(gf['education_training'])}</td></tr>
              <tr><td class="lab">Опыт практической работы</td>
                  <td>{esc(gf['practical_experience'] or '-')}</td></tr>
            </table>
            <table>
              <tr><td class="lab">Особые условия допуска к работе</td>
                  <td>{esc(gf['special_admission'] or '-')}</td></tr>
              <tr><td class="lab">Другие характеристики</td>
                  <td>{esc(gf['other_characteristics'] or '-')}</td></tr>
            </table>
            <p class="sub">Справочная информация</p>
            <table>
              <tr><th>Наименование документа</th><th>Код</th>
                  <th>Наименование базовой группы, должности (профессии) или специальности</th></tr>
              {''.join(ref_rows)}
            </table>
            {''.join(pf_parts)}
            """
        )

    co = "".join(
        f"<tr><td class='num'>{n}</td><td>{esc(o)}</td></tr>"
        for n, o in enumerate(doc["co_developers"] or [""], start=1)
    )
    abbrs = []
    for a in doc["abbreviations"] or []:
        if isinstance(a, dict):
            if a.get("abbr"):
                abbrs.append(f"{esc(a['abbr'])} – {esc(a.get('meaning') or '')}")
            elif a.get("text"):
                abbrs.append(esc(a["text"]))

    head = doc["developer_head"] or ""
    head_pos, head_fio = head, ""
    m = re.match(
        r"^(Президент|Генеральный директор|Директор|Председатель|Руководитель)\s+(.+)$",
        head,
        re.I,
    )
    if m:
        head_pos, head_fio = m.group(1), m.group(2)

    return f"""<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8"/>
<title>ПС {esc(doc['reg_number'])} — {esc(doc['name'])}</title>
<style>
@page {{ margin: 2cm 1cm 2cm 2cm; }}
body{{font-family:"Times New Roman",Times,serif;font-size:12pt;line-height:1.25;
      margin:2cm 1cm 2cm 2cm;color:#000}}
h1{{text-align:center;font-size:16pt;margin:0.6em 0}}
h2{{font-size:13pt;margin:1em 0 0.4em}}
h3{{font-size:12pt;margin:0.8em 0 0.3em}}
.stamp{{text-align:right;font-size:10pt;line-height:1.25;margin:0 0 12px 0;border:none}}
.name{{text-align:center;font-size:14pt;font-weight:bold;margin:0.4em 0 0.8em}}
.toc{{margin:1em 0 1.5em;line-height:1.5}}
.toc-title{{text-align:center;font-weight:bold;margin-bottom:0.4em}}
.sub{{font-weight:bold;margin:0.6em 0 0.25em}}
.sec-i{{page-break-before:always}}
table{{border-collapse:collapse;width:100%;margin:0.3em 0 0.7em;table-layout:fixed}}
td,th{{border:1px solid #000;padding:3px 5px;vertical-align:top}}
th{{text-align:center;font-weight:bold;font-size:10pt}}
td.lab{{font-weight:bold;width:22%}}
td.num{{width:1.5em;text-align:center}}
tr.cap td{{text-align:center;font-size:9pt}}
.reg-wrap{{display:flex;justify-content:flex-end;margin:0.8em 0}}
.reg{{width:6.5cm;text-align:center}}
.meta6 td.lab{{width:12%;white-space:nowrap}}
.toolbar{{margin-bottom:10px}}
@media print{{.toolbar{{display:none}}}}
</style></head><body>
<div class="toolbar"><button onclick="window.print()">Печать</button></div>
<div class="stamp">{stamp_html}</div>
<h1>ПРОФЕССИОНАЛЬНЫЙ СТАНДАРТ</h1>
<div class="name">{esc(doc['name'])}</div>
<div class="reg-wrap"><table class="reg">
<tr><td style="font-size:14pt;padding:10px">{esc(doc['reg_number'])}</td></tr>
<tr><td>Регистрационный номер</td></tr>
</table></div>
<div class="toc"><div class="toc-title">Содержание</div>{toc_html}</div>

<div class="sec-i">
<h2>I. Общие сведения</h2>
<table>
<tr><td style="width:80%">{esc(doc['kind_activity'])}</td>
    <td style="width:20%;text-align:center">{esc(doc['ps_code'])}</td></tr>
<tr class="cap"><td>(наименование вида профессиональной деятельности)</td><td>Код</td></tr>
</table>
<p class="sub">Краткое описание вида профессиональной деятельности</p>
<table><tr><td>{esc(doc['purpose'])}</td></tr></table>
<p class="sub">Группа занятий:</p>
<table>{''.join(okz_rows_html)}</table>
<p class="sub">Отнесение к области профессиональной деятельности:</p>
<table>
<tr><td style="width:15%">{esc(doc['opd_code'])}</td><td>{esc(doc['opd_name'])}</td></tr>
<tr class="cap"><td>(код ОПД)</td><td>(наименование области профессиональной деятельности)</td></tr>
</table>
<p class="sub">Отнесение к видам экономической деятельности:</p>
<table>{okved_body}</table>
</div>

<h2>II. Описание трудовых функций, входящих в профессиональный стандарт
(функциональная карта вида профессиональной деятельности)</h2>
<table>
<tr><th colspan="4">Обобщенные трудовые функции</th><th colspan="3">Трудовые функции</th></tr>
<tr>
<th>код</th><th>наименование</th><th>уровень квалификации</th>
<th>Возможные наименования должностей, профессий рабочих</th>
<th>наименование</th><th>код</th><th>уровень (подуровень) квалификации</th>
</tr>
{map_html}
</table>

<h2>III. Характеристика обобщенных трудовых функций</h2>
{''.join(gf_html)}

<h2>IV. Сведения об организациях – разработчиках профессионального стандарта</h2>
<h3>4.1. Ответственная организация-разработчик</h3>
<table>
<tr><td colspan="2">{esc(doc['developer_org'])}</td></tr>
<tr><td>{esc(head_pos)}</td><td>{esc(head_fio)}</td></tr>
</table>
<h3>4.2. Наименования организаций-разработчиков</h3>
<table>{co}</table>

<h2>V. Сокращения, используемые в профессиональном стандарте</h2>
<p>{'<br/>'.join(abbrs) if abbrs else ''}</p>
</body></html>
"""


def build_ps_docx_bytes(doc: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    document = Document()
    for section in document.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.0)

    def font(run, size=11, bold=False):
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), "Times New Roman")

    def p(text="", *, bold=False, size=11, center=False, right=False, space_after=6):
        para = document.add_paragraph()
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif right:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(0)
        if text:
            run = para.add_run(text)
            font(run, size=size, bold=bold)
        return para

    def cell(c, text, *, bold=False, size=10, center=False, vcenter=False):
        c.text = ""
        para = c.paragraphs[0]
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text or "")
        font(run, size=size, bold=bold)
        if vcenter:
            tc = c._tc
            tcPr = tc.get_or_add_tcPr()
            vAlign = OxmlElement("w:vAlign")
            vAlign.set(qn("w:val"), "center")
            tcPr.append(vAlign)

    def table(rows, cols):
        t = document.add_table(rows=rows, cols=cols)
        t.style = "Table Grid"
        t.autofit = False
        t.allow_autofit = False
        return t

    def set_widths(t, widths_cm: list[float]):
        total = sum(widths_cm)
        if total > 0 and abs(total - CONTENT_W) > 0.05:
            widths_cm = [w * CONTENT_W / total for w in widths_cm]
        tbl = t._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(round(CONTENT_W * 567))))
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)

    def page_break():
        para = document.add_paragraph()
        run = para.add_run()
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run._r.append(br)

    # 1) Штамп — абзацы справа, без таблицы/рамки; дата и № из БД
    for i, line in enumerate(_format_order_stamp(doc["approval_date"], doc["order_number"])):
        p(line, size=10, right=True, bold=(i == 0), space_after=0)
    p("", space_after=10)

    p("ПРОФЕССИОНАЛЬНЫЙ СТАНДАРТ", bold=True, size=16, center=True, space_after=8)
    p(doc["name"], bold=True, size=14, center=True, space_after=12)

    # 2) Рег. номер — таблица у правого края
    rt = table(2, 1)
    rt.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for row in rt.rows:
        row.cells[0].width = Cm(6.5)
    cell(rt.cell(0, 0), doc["reg_number"], size=14, center=True)
    cell(rt.cell(1, 0), "Регистрационный номер", size=10, center=True)
    p("", space_after=6)

    p("Содержание", bold=True, size=12, center=True)
    toc_items = [
        "I. Общие сведения",
        "II. Описание трудовых функций, входящих в профессиональный стандарт "
        "(функциональная карта вида профессиональной деятельности)",
        "III. Характеристика обобщенных трудовых функций",
    ]
    for i, gf in enumerate(doc["generalized_functions"], start=1):
        toc_items.append(f"3.{i}. Обобщенная трудовая функция «{gf['name']}»")
    toc_items += [
        "IV. Сведения об организациях – разработчиках профессионального стандарта",
        "4.1. Ответственная организация-разработчик",
        "4.2. Наименования организаций-разработчиков",
        "V. Сокращения, используемые в профессиональном стандарте",
    ]
    for item in toc_items:
        p(item, size=11, space_after=2)

    # 3) Раздел I — с новой страницы
    page_break()
    p("I. Общие сведения", bold=True, size=13, center=True, space_after=8)

    # 4) Все таблицы раздела I — одинаковая ширина CONTENT_W
    t = table(2, 2)
    set_widths(t, [14.5, 3.5])
    cell(t.cell(0, 0), doc["kind_activity"], size=11)
    cell(t.cell(0, 1), doc["ps_code"], size=11, center=True, vcenter=True)
    cell(t.cell(1, 0), "(наименование вида профессиональной деятельности)", size=8, center=True)
    cell(t.cell(1, 1), "Код", size=8, center=True)

    p("Краткое описание вида профессиональной деятельности", bold=True, size=11)
    t = table(1, 1)
    set_widths(t, [CONTENT_W])
    cell(t.cell(0, 0), doc["purpose"], size=11)

    p("Группа занятий:", bold=True, size=11)
    pairs = _split_okz_pairs(doc["okz_group_code"], doc["okz_group_name"])
    n_pair_rows = max(1, (len(pairs) + 1) // 2)
    t = table(n_pair_rows + 1, 4)
    set_widths(t, [2.5, 6.5, 2.5, 6.5])
    for i in range(n_pair_rows):
        left = pairs[i * 2] if i * 2 < len(pairs) else ("", "")
        right = pairs[i * 2 + 1] if i * 2 + 1 < len(pairs) else ("", "")
        cell(t.cell(i, 0), left[0], size=10)
        cell(t.cell(i, 1), left[1], size=10)
        cell(t.cell(i, 2), right[0], size=10)
        cell(t.cell(i, 3), right[1], size=10)
    for c, lab in enumerate(["(код ОКЗ)", "(наименование)", "(код ОКЗ)", "(наименование)"]):
        cell(t.cell(n_pair_rows, c), lab, size=8, center=True)

    p("Отнесение к области профессиональной деятельности:", bold=True, size=11)
    t = table(2, 2)
    set_widths(t, [2.5, 15.5])
    cell(t.cell(0, 0), doc["opd_code"], size=10)
    cell(t.cell(0, 1), doc["opd_name"], size=10)
    cell(t.cell(1, 0), "(код ОПД)", size=8, center=True)
    cell(t.cell(1, 1), "(наименование области профессиональной деятельности)", size=8, center=True)

    p("Отнесение к видам экономической деятельности:", bold=True, size=11)
    units = doc["okved_units"] or [{"code": "", "name": ""}]
    t = table(len(units) + 1, 2)
    set_widths(t, [2.5, 15.5])
    for i, u in enumerate(units):
        cell(t.cell(i, 0), u.get("code") or "", size=10)
        cell(t.cell(i, 1), u.get("name") or "", size=10)
    cell(t.cell(len(units), 0), "(код ОКВЭД)", size=8, center=True)
    cell(t.cell(len(units), 1), "(наименование вида экономической деятельности)", size=8, center=True)

    # 6) Функциональная карта — строки ТФ + вертикальное объединение ОТФ
    p(
        "II. Описание трудовых функций, входящих в профессиональный стандарт "
        "(функциональная карта вида профессиональной деятельности)",
        bold=True,
        size=12,
        center=True,
        space_after=8,
    )

    map_rows = _build_map_rows(doc)
    if not map_rows:
        map_rows = [
            {
                "otf_code": "",
                "otf_name": "",
                "otf_level": "",
                "titles": "",
                "tf_name": "",
                "tf_code": "",
                "tf_level": "",
            }
        ]

    t = table(2 + len(map_rows), 7)
    set_widths(t, [1.2, 4.0, 1.5, 3.5, 4.0, 1.8, 2.0])

    a = t.cell(0, 0)
    a.merge(t.cell(0, 3))
    cell(a, "Обобщенные трудовые функции", bold=True, size=9, center=True)
    b = t.cell(0, 4)
    b.merge(t.cell(0, 6))
    cell(b, "Трудовые функции", bold=True, size=9, center=True)

    headers = [
        "код",
        "наименование",
        "уровень квалификации",
        "Возможные наименования должностей, профессий рабочих",
        "наименование",
        "код",
        "уровень (подуровень) квалификации",
    ]
    for c, h in enumerate(headers):
        cell(t.cell(1, c), h, bold=True, size=8, center=True)

    for r, row in enumerate(map_rows):
        rr = 2 + r
        cell(t.cell(rr, 0), row["otf_code"], size=8, center=True, vcenter=True)
        cell(t.cell(rr, 1), row["otf_name"], size=8, vcenter=True)
        cell(t.cell(rr, 2), row["otf_level"], size=8, center=True, vcenter=True)
        cell(t.cell(rr, 3), row["titles"], size=8, vcenter=True)
        cell(t.cell(rr, 4), row["tf_name"], size=8)
        cell(t.cell(rr, 5), row["tf_code"], size=8, center=True)
        cell(t.cell(rr, 6), row["tf_level"], size=8, center=True)

    # Вертикальный merge колонок ОТФ
    block_start = 0
    for idx in range(1, len(map_rows) + 1):
        cont = (
            idx < len(map_rows)
            and map_rows[idx]["otf_code"] == map_rows[block_start]["otf_code"]
            and map_rows[idx]["otf_name"] == map_rows[block_start]["otf_name"]
        )
        if not cont:
            end = idx - 1
            if end > block_start:
                for col in range(4):
                    t.cell(2 + block_start, col).merge(t.cell(2 + end, col))
            block_start = idx

    # ========== III ==========
    p("III. Характеристика обобщенных трудовых функций", bold=True, size=13, center=True)

    for i, gf in enumerate(doc["generalized_functions"], start=1):
        p(
            f"3.{i}. Обобщенная трудовая функция «{gf['name']}»",
            bold=True,
            size=12,
            space_after=6,
        )

        t = table(1, 6)
        set_widths(t, [2.8, 7.5, 1.0, 1.5, 2.5, 2.7])
        cell(t.cell(0, 0), "Наименование", bold=True, size=9)
        cell(t.cell(0, 1), gf["name"], size=10)
        cell(t.cell(0, 2), "Код", bold=True, size=9)
        cell(t.cell(0, 3), gf["code"], size=10, center=True)
        cell(t.cell(0, 4), "Уровень квалификации", bold=True, size=9)
        cell(t.cell(0, 5), gf["level"], size=10, center=True)

        t = table(1, 2)
        set_widths(t, [3.5, 14.5])
        cell(t.cell(0, 0), "Возможные наименования должностей, профессий рабочих", bold=True, size=9)
        cell(t.cell(0, 1), _job_titles_text(gf["possible_job_titles"]), size=10)

        p("Пути достижения квалификации", bold=True, size=11)
        t = table(2, 2)
        set_widths(t, [3.5, 14.5])
        cell(t.cell(0, 0), "Образование и обучение", bold=True, size=9)
        cell(t.cell(0, 1), gf["education_training"] or "", size=10)
        cell(t.cell(1, 0), "Опыт практической работы", bold=True, size=9)
        cell(t.cell(1, 1), gf["practical_experience"] or "-", size=10)

        t = table(2, 2)
        set_widths(t, [3.5, 14.5])
        cell(t.cell(0, 0), "Особые условия допуска к работе", bold=True, size=9)
        cell(t.cell(0, 1), gf["special_admission"] or "-", size=10)
        cell(t.cell(1, 0), "Другие характеристики", bold=True, size=9)
        cell(t.cell(1, 1), gf["other_characteristics"] or "-", size=10)

        p("Справочная информация", bold=True, size=11)
        ref_rows: list[list[str]] = []
        for label, key in REF_LABELS:
            units = gf.get(key) or []
            if not units:
                ref_rows.append([label, "", ""])
            else:
                for u in units:
                    ref_rows.append([label, u.get("code") or "", u.get("name") or ""])
        t = table(1 + len(ref_rows), 3)
        set_widths(t, [3.5, 3.0, 11.5])
        cell(t.cell(0, 0), "Наименование документа", bold=True, size=8, center=True)
        cell(t.cell(0, 1), "Код", bold=True, size=8, center=True)
        cell(
            t.cell(0, 2),
            "Наименование базовой группы, должности (профессии) или специальности",
            bold=True,
            size=8,
            center=True,
        )
        for r, row in enumerate(ref_rows, start=1):
            for c, val in enumerate(row):
                cell(t.cell(r, c), val, size=9)

        for j, pf in enumerate(gf["particular_functions"], start=1):
            p(f"3.{i}.{j}. Трудовая функция", bold=True, size=12, space_after=4)
            t = table(1, 6)
            set_widths(t, [2.8, 7.2, 1.0, 1.8, 2.6, 2.6])
            cell(t.cell(0, 0), "Наименование", bold=True, size=9)
            cell(t.cell(0, 1), pf["name"], size=10)
            cell(t.cell(0, 2), "Код", bold=True, size=9)
            cell(t.cell(0, 3), pf["code"], size=10, center=True)
            cell(t.cell(0, 4), "Уровень (подуровень) квалификации", bold=True, size=9)
            cell(t.cell(0, 5), pf["sub_qualification"], size=10, center=True)

            body_rows: list[tuple[str, str]] = []
            for item in pf["labor_actions"] or []:
                body_rows.append(("Трудовые действия", item))
            if not pf["labor_actions"]:
                body_rows.append(("Трудовые действия", ""))
            if pf["required_skills"]:
                for item in pf["required_skills"]:
                    body_rows.append(("Необходимые умения", item))
            else:
                body_rows.append(("Необходимые умения", ""))
            if pf["necessary_knowledges"]:
                for item in pf["necessary_knowledges"]:
                    body_rows.append(("Необходимые знания", item))
            else:
                body_rows.append(("Необходимые знания", ""))
            body_rows.append(("Другие характеристики", pf["other_characteristics"] or "-"))

            t = table(len(body_rows), 2)
            set_widths(t, [3.6, 14.4])
            for r, (lab, val) in enumerate(body_rows):
                cell(t.cell(r, 0), lab, bold=True, size=9)
                cell(t.cell(r, 1), val, size=10)

    p(
        "IV. Сведения об организациях – разработчиках профессионального стандарта",
        bold=True,
        size=12,
        center=True,
    )
    p("4.1. Ответственная организация-разработчик", bold=True, size=12)
    head = doc["developer_head"] or ""
    head_pos, head_fio = head, ""
    m = re.match(
        r"^(Президент|Генеральный директор|Директор|Председатель|Руководитель)\s+(.+)$",
        head,
        re.I,
    )
    if m:
        head_pos, head_fio = m.group(1), m.group(2)
    t = table(2, 2)
    set_widths(t, [9.0, 9.0])
    t.cell(0, 0).merge(t.cell(0, 1))
    cell(t.cell(0, 0), doc["developer_org"], size=10)
    cell(t.cell(1, 0), head_pos, size=10)
    cell(t.cell(1, 1), head_fio, size=10)

    p("4.2. Наименования организаций-разработчиков", bold=True, size=12)
    orgs = doc["co_developers"] or [""]
    t = table(len(orgs), 2)
    set_widths(t, [1.2, 16.8])
    for n, org in enumerate(orgs):
        cell(t.cell(n, 0), str(n + 1), size=10, center=True)
        cell(t.cell(n, 1), org or "", size=10)

    p("V. Сокращения, используемые в профессиональном стандарте", bold=True, size=12, center=True)
    if not doc["abbreviations"]:
        p("")
    else:
        for a in doc["abbreviations"]:
            if isinstance(a, dict):
                if a.get("abbr"):
                    p(f"{a['abbr']} – {a.get('meaning') or ''}", size=11, space_after=2)
                elif a.get("text"):
                    p(a["text"], size=11, space_after=2)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
