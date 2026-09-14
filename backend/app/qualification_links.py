"""Связь квалификаций НАРК с профессиональными стандартами."""
from __future__ import annotations

import os
import re
from collections import defaultdict
from pathlib import Path
from typing import NamedTuple

from sqlalchemy import func
from sqlalchemy.orm import Session

from .db.qualifications_models import Qualification
from .db.raw_models import StandardRaw
from .db.assessment_tools_models import AssessmentTool

# Код квалификации: 01.00100.01 / 40.20900.100 → код ПС 01.001 / 40.209
QUALIFICATION_CODE_RE = re.compile(r"^(\d{2}\.\d{3})\d{2}\.\d{2,4}$")
PS_CODE_RE = re.compile(r"^(\d{2})\.(\d{1,5})$")
ORDER_TOKEN_RE = re.compile(r"(\d{1,4}\s*[а-яa-z])", re.IGNORECASE)
NO_LINKED_PS = "Нет связанного профессионального стандарта"


class LinkResult(NamedTuple):
    linked: int
    skipped: int
    not_found: int
    by_method: dict[str, int]


def ps_code_from_qualification_code(code: str) -> str | None:
    """01.00100.01 → 01.001; 03.01400.02 → 03.014."""
    match = QUALIFICATION_CODE_RE.match((code or "").strip())
    if match:
        return match.group(1)
    return None


def normalize_ps_code(code: str) -> str | None:
    text = (code or "").strip()
    if not text:
        return None
    parts = text.split(".")
    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
        return f"{int(parts[0]):02d}.{int(parts[1]):03d}"
    match = PS_CODE_RE.match(text)
    if not match:
        return None
    area, seq = match.groups()
    return f"{area}.{int(seq):03d}"


def normalize_name(text: str) -> str:
    value = (text or "").strip().lower().replace("ё", "е")
    value = re.sub(r"[«»\"'`]", "", value)
    # Унификация дефисов/тире: «Врач – травматолог» ≈ «Врач-травматолог»
    value = re.sub(r"[\s]*[‐‑‒–—―−-][\s]*", "-", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" .,-—–;")


def normalize_order_number(order_text: str) -> str:
    if not order_text:
        return ""
    match = re.search(r"(?:N|№|n|N\.?)\s*(\d+[а-яa-z]?)", order_text, re.IGNORECASE)
    if match:
        return match.group(1).strip().lower()
    parts = order_text.split()
    last = parts[-1] if parts else ""
    if re.match(r"\d+[а-яa-z]?", last, re.IGNORECASE):
        return last.lower()
    return ""


def extract_order_tokens(order_text: str) -> set[str]:
    tokens: set[str] = set()
    normalized = normalize_order_number(order_text)
    if normalized:
        tokens.add(normalized)
    for match in ORDER_TOKEN_RE.finditer(order_text or ""):
        token = re.sub(r"\s+", "", match.group(1)).lower()
        if token:
            tokens.add(token)
    return tokens


def extract_order_year(order_text: str) -> str | None:
    """Год из реквизитов приказа: «от 11.10.2021 № 696н» → 2021."""
    text = order_text or ""
    match = re.search(r"от\s+\d{1,2}[./]\d{1,2}[./](\d{4})", text, re.IGNORECASE)
    if match:
        return match.group(1)
    match = re.search(r"(20\d{2})", text)
    return match.group(1) if match else None


def standard_ps_codes(std: StandardRaw) -> set[str]:
    codes: set[str] = set()
    if std.ps_code:
        normalized = normalize_ps_code(std.ps_code)
        if normalized:
            codes.add(normalized)
    element_id = std.element_id or ""
    if element_id.startswith("classinform:"):
        normalized = normalize_ps_code(element_id.split(":", 1)[1])
        if normalized:
            codes.add(normalized)
    return codes


def names_match(standard_name: str, qualification_name: str) -> bool:
    left = normalize_name(standard_name)
    right = normalize_name(qualification_name)
    if not left or not right:
        return False
    if left == right:
        return True
    if len(right) >= 15 and right in left:
        return True
    if len(left) >= 15 and left in right:
        return True
    if len(left) >= 20 and len(right) >= 20 and left[:40] == right[:40]:
        return True
    return False


def _filter_candidates_compatible(
    candidates: list[StandardRaw],
    qualification: Qualification,
    *,
    require_name: bool = False,
    ignore_code_conflict: bool = False,
) -> list[StandardRaw]:
    """Отсекает ложные совпадения: чужой код ПС, другой год приказа, без совпадения названия."""
    if not candidates:
        return []

    qual_ps_code = ps_code_from_qualification_code(qualification.code or "")
    qual_name = (qualification.prof_standard_name or "").strip()
    qual_year = extract_order_year(qualification.prof_standard_order or "")

    filtered: list[StandardRaw] = []
    for std in candidates:
        std_codes = standard_ps_codes(std)
        # Квалификация 40.071* не должна цепляться к ПС 12.013 только из‑за № приказа
        if (
            not ignore_code_conflict
            and qual_ps_code
            and std_codes
            and qual_ps_code not in std_codes
        ):
            continue

        if require_name:
            if not qual_name or qual_name == NO_LINKED_PS:
                continue
            if not names_match(std.name or "", qual_name):
                continue

        if qual_year:
            std_year = extract_order_year(std.order_number or "") or extract_order_year(
                std.approval_date or ""
            )
            if std_year and std_year != qual_year:
                continue

        filtered.append(std)

    return filtered


class StandardLookup:
    """Индексы профстандартов для быстрого сопоставления."""

    def __init__(self, session: Session):
        self.by_ps_code: dict[str, list[StandardRaw]] = defaultdict(list)
        self.by_order_token: dict[str, list[StandardRaw]] = defaultdict(list)
        self.all: list[StandardRaw] = session.query(StandardRaw).all()
        for std in self.all:
            for code in standard_ps_codes(std):
                self.by_ps_code[code].append(std)
            for token in extract_order_tokens(std.order_number or ""):
                self.by_order_token[token].append(std)

    def candidates_for_ps_code(self, ps_code: str) -> list[StandardRaw]:
        normalized = normalize_ps_code(ps_code) or ps_code
        return list(self.by_ps_code.get(normalized, []))


def pick_best_standard(candidates: list[StandardRaw], qualification: Qualification) -> StandardRaw | None:
    if not candidates:
        return None
    if len(candidates) == 1:
        return candidates[0]

    active = [
        std for std in candidates
        if (getattr(std, "status", None) or "active") == "active"
    ]
    pool = active or candidates

    order_norm = normalize_order_number(qualification.prof_standard_order or "")
    qual_year = extract_order_year(qualification.prof_standard_order or "")
    if order_norm:
        year_and_order: list[StandardRaw] = []
        for std in pool:
            order_tokens = extract_order_tokens(std.order_number or "")
            if order_norm not in order_tokens and order_norm not in (std.order_number or "").lower():
                continue
            if qual_year:
                std_year = extract_order_year(std.order_number or "") or extract_order_year(
                    std.approval_date or ""
                )
                if std_year and std_year != qual_year:
                    continue
            year_and_order.append(std)
        if year_and_order:
            pool = year_and_order

    qual_name = qualification.prof_standard_name or ""
    if qual_name and qual_name != NO_LINKED_PS:
        named = [std for std in pool if names_match(std.name or "", qual_name)]
        if named:
            return named[0]

    return pool[0]


def find_standard_for_qualification(
    qualification: Qualification,
    lookup: StandardLookup,
) -> tuple[StandardRaw | None, str | None]:
    ps_code = ps_code_from_qualification_code(qualification.code or "")
    if ps_code:
        candidates = lookup.candidates_for_ps_code(ps_code)
        chosen = pick_best_standard(candidates, qualification)
        if chosen:
            return chosen, f"код квалификации ({ps_code})"
        # Код из квалификации (напр. 10.104) может не совпадать с кодом ПС (40.104).
        # Тогда пробуем название/приказ — но НЕ чужой код через номер приказа вслепую.

    qual_name = (qualification.prof_standard_name or "").strip()
    if qual_name and qual_name != NO_LINKED_PS:
        name_hits = [
            std for std in lookup.all
            if names_match(std.name or "", qual_name)
        ]
        # При известном коде из квалификации не отсекаем по конфликту кода —
        # иначе 10.104* никогда не свяжется с ПС 40.104.
        name_hits = _filter_candidates_compatible(
            name_hits,
            qualification,
            require_name=False,
            ignore_code_conflict=bool(ps_code),
        )
        chosen = pick_best_standard(name_hits, qualification)
        if chosen:
            return chosen, "точное/частичное название"

    order_norm = normalize_order_number(qualification.prof_standard_order or "")
    if order_norm:
        candidates = lookup.by_order_token.get(order_norm, [])
        # Номер приказа без названия/года часто совпадает у разных ПС (напр. 696н 2020 и 2021)
        candidates = _filter_candidates_compatible(
            candidates, qualification, require_name=True
        )
        chosen = pick_best_standard(candidates, qualification)
        if chosen:
            return chosen, "номер приказа+название"

    # Частичный приказ — только вместе с совпадением названия и без конфликта кода ПС
    if qualification.prof_standard_order and qual_name and qual_name != NO_LINKED_PS:
        order_clean = re.sub(r"[^0-9а-яa-z]", "", qualification.prof_standard_order.lower())
        if len(order_clean) > 3:
            partial: list[StandardRaw] = []
            for std in lookup.all:
                std_order = re.sub(r"[^0-9а-яa-z]", "", (std.order_number or "").lower())
                if order_clean in std_order and names_match(std.name or "", qual_name):
                    partial.append(std)
            partial = _filter_candidates_compatible(partial, qualification, require_name=True)
            chosen = pick_best_standard(partial, qualification)
            if chosen:
                return chosen, "частичный приказ+название"

    return None, None


def link_qualification_record(
    qualification: Qualification,
    lookup: StandardLookup,
    *,
    update_existing: bool = True,
) -> tuple[bool, str | None]:
    if qualification.prof_standard_id is not None and not update_existing:
        return False, None

    standard, method = find_standard_for_qualification(qualification, lookup)
    if standard:
        qualification.prof_standard_id = standard.id
        return True, method

    if update_existing:
        qualification.prof_standard_id = None
    return False, None


def _registry_xlsx_path() -> str | None:
    path_file = Path(__file__).resolve().parent.parent / "scripts" / "xlsx_path.txt"
    if not path_file.is_file():
        return None
    text = path_file.read_text(encoding="utf-8-sig").strip()
    return text if text and os.path.isfile(text) else None


def backfill_ps_codes_from_registry_xlsx(session: Session) -> int:
    """Заполняет ps_code у ПС из XLSX реестра Минтруда (если файл указан в scripts/xlsx_path.txt)."""
    xlsx_path = _registry_xlsx_path()
    if not xlsx_path:
        return 0

    try:
        from openpyxl import load_workbook
    except ImportError:
        return 0

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return 0

    header = [str(cell or "").strip().lower() for cell in rows[0]]

    def col(*needles: str) -> int | None:
        for index, name in enumerate(header):
            if any(needle in name for needle in needles):
                return index
        return None

    reg_idx = col("регистрац", "рег")
    code_idx = col("код")
    if reg_idx is None:
        return 0

    by_reg: dict[str, str] = {}
    for row in rows[1:]:
        if not row or reg_idx >= len(row):
            continue
        reg = str(row[reg_idx] or "").strip()
        if not reg:
            continue
        ps_code = ""
        if code_idx is not None and code_idx < len(row):
            ps_code = normalize_ps_code(str(row[code_idx] or "").strip()) or ""
        if ps_code:
            by_reg[reg] = ps_code

    updated = 0
    for std in session.query(StandardRaw).filter(StandardRaw.ps_code.is_(None)).all():
        ps_code = by_reg.get(std.reg_number)
        if not ps_code:
            continue
        std.ps_code = ps_code
        if not std.professional_area_code:
            std.professional_area_code = ps_code.split(".")[0]
        updated += 1

    if updated:
        session.flush()
    return updated


def relink_all_qualifications(
    session: Session,
    *,
    update_existing: bool = True,
    commit: bool = True,
    backfill_ps_codes: bool = True,
) -> LinkResult:
    if backfill_ps_codes:
        filled = backfill_ps_codes_from_registry_xlsx(session)
        if filled:
            session.flush()

    lookup = StandardLookup(session)
    qualifications = session.query(Qualification).all()

    linked = 0
    skipped = 0
    not_found = 0
    by_method: dict[str, int] = defaultdict(int)

    for qualification in qualifications:
        if qualification.prof_standard_id is not None and not update_existing:
            skipped += 1
            continue

        was_linked, method = link_qualification_record(
            qualification,
            lookup,
            update_existing=update_existing,
        )
        if was_linked and method:
            linked += 1
            by_method[method] += 1
        elif qualification.prof_standard_id is None:
            not_found += 1
        else:
            skipped += 1

    if commit:
        session.commit()

    return LinkResult(
        linked=linked,
        skipped=skipped,
        not_found=not_found,
        by_method=dict(by_method),
    )


def qualification_count_by_standard_id(session: Session) -> dict[int, int]:
    rows = (
        session.query(Qualification.prof_standard_id, func.count(Qualification.id))
        .filter(Qualification.prof_standard_id.isnot(None))
        .group_by(Qualification.prof_standard_id)
        .all()
    )
    return {int(standard_id): int(count) for standard_id, count in rows}


def assessment_tool_count_by_qualification_id(session: Session) -> dict[int, int]:
    rows = (
        session.query(AssessmentTool.qualification_id, func.count(AssessmentTool.id))
        .filter(AssessmentTool.qualification_id.isnot(None))
        .filter(
            (AssessmentTool.status == "active")
            | (AssessmentTool.status.is_(None))
            | (AssessmentTool.status == "")
        )
        .group_by(AssessmentTool.qualification_id)
        .all()
    )
    return {int(qid): int(count) for qid, count in rows}


def assessment_tool_count_by_standard_id(session: Session) -> dict[int, int]:
    """Сумма активных ОС по квалификациям, привязанным к ПС."""
    rows = (
        session.query(Qualification.prof_standard_id, func.count(AssessmentTool.id))
        .join(AssessmentTool, AssessmentTool.qualification_id == Qualification.id)
        .filter(Qualification.prof_standard_id.isnot(None))
        .filter(
            (AssessmentTool.status == "active")
            | (AssessmentTool.status.is_(None))
            | (AssessmentTool.status == "")
        )
        .group_by(Qualification.prof_standard_id)
        .all()
    )
    return {int(standard_id): int(count) for standard_id, count in rows}


def qualification_flags(standard_id: int, counts: dict[int, int]) -> dict[str, int | bool]:
    count = counts.get(standard_id, 0)
    return {
        "qualification_count": count,
        "has_qualifications": count > 0,
    }


def assessment_tool_flags(owner_id: int, counts: dict[int, int]) -> dict[str, int | bool]:
    count = counts.get(owner_id, 0)
    return {
        "assessment_tool_count": count,
        "has_assessment_tools": count > 0,
    }


def standards_qualification_coverage(session: Session) -> dict[str, int]:
    counts = qualification_count_by_standard_id(session)
    total_standards = session.query(StandardRaw).count()
    total_qualifications = session.query(Qualification).count()
    linked_qualifications = session.query(Qualification).filter(
        Qualification.prof_standard_id.isnot(None)
    ).count()
    return {
        "total_standards": total_standards,
        "standards_with_qualifications": len(counts),
        "standards_without_qualifications": max(0, total_standards - len(counts)),
        "total_qualifications": total_qualifications,
        "linked_qualifications": linked_qualifications,
        "unlinked_qualifications": max(0, total_qualifications - linked_qualifications),
    }


def audit_qualification_links(session: Session) -> dict:
    """Проверка текущих связей квалификация ↔ ПС на типичные ошибки."""
    standards = {std.id: std for std in session.query(StandardRaw).all()}
    qualifications = session.query(Qualification).all()

    code_mismatch: list[dict] = []
    name_mismatch: list[dict] = []
    year_mismatch: list[dict] = []
    missing_target_ps: list[dict] = []
    unlinked_with_known_code: list[dict] = []

    lookup = StandardLookup(session)
    order_collisions: dict[str, list[dict]] = defaultdict(list)
    for std in standards.values():
        for token in extract_order_tokens(std.order_number or ""):
            year = extract_order_year(std.order_number or "") or extract_order_year(
                std.approval_date or ""
            )
            order_collisions[token].append(
                {
                    "id": std.id,
                    "reg_number": std.reg_number,
                    "ps_code": std.ps_code,
                    "name": std.name,
                    "year": year,
                    "order_number": std.order_number,
                }
            )

    colliding_orders = {
        token: items
        for token, items in order_collisions.items()
        if len({(i.get("year"), i.get("ps_code"), i.get("id")) for i in items}) > 1
        and len(items) > 1
    }

    for qual in qualifications:
        qual_code = ps_code_from_qualification_code(qual.code or "")
        std = standards.get(qual.prof_standard_id) if qual.prof_standard_id else None

        if qual_code and qual.prof_standard_id is None:
            has_target = bool(lookup.candidates_for_ps_code(qual_code))
            row = {
                "qualification_id": qual.id,
                "code": qual.code,
                "name": qual.name,
                "expected_ps_code": qual_code,
                "ps_in_db": has_target,
                "prof_standard_name": qual.prof_standard_name,
            }
            if has_target:
                unlinked_with_known_code.append(row)
            else:
                missing_target_ps.append(row)
            continue

        if not std:
            continue

        std_codes = standard_ps_codes(std)
        code_ok = bool(qual_code and std_codes and qual_code in std_codes)
        if qual_code and std_codes and not code_ok:
            code_mismatch.append(
                {
                    "qualification_id": qual.id,
                    "code": qual.code,
                    "name": qual.name,
                    "expected_ps_code": qual_code,
                    "linked_standard_id": std.id,
                    "linked_ps_code": std.ps_code,
                    "linked_standard_name": std.name,
                    "linked_reg_number": std.reg_number,
                }
            )

        qual_ps_name = (qual.prof_standard_name or "").strip()
        if (
            qual_ps_name
            and qual_ps_name != NO_LINKED_PS
            and std.name
            and not names_match(std.name, qual_ps_name)
            and not code_ok
        ):
            name_mismatch.append(
                {
                    "qualification_id": qual.id,
                    "code": qual.code,
                    "name": qual.name,
                    "qualification_ps_name": qual_ps_name,
                    "linked_standard_name": std.name,
                    "linked_standard_id": std.id,
                    "linked_ps_code": std.ps_code,
                }
            )

        qual_year = extract_order_year(qual.prof_standard_order or "")
        std_year = extract_order_year(std.order_number or "") or extract_order_year(
            std.approval_date or ""
        )
        # Расхождение года при совпадающем коде — обычно смена редакции ПС, не ложная связь
        if qual_year and std_year and qual_year != std_year and not code_ok:
            year_mismatch.append(
                {
                    "qualification_id": qual.id,
                    "code": qual.code,
                    "name": qual.name,
                    "qualification_order": qual.prof_standard_order,
                    "qualification_year": qual_year,
                    "linked_standard_id": std.id,
                    "linked_order": std.order_number,
                    "linked_year": std_year,
                    "linked_standard_name": std.name,
                }
            )

    coverage = standards_qualification_coverage(session)
    return {
        **coverage,
        "issues": {
            "code_mismatch": len(code_mismatch),
            "name_mismatch": len(name_mismatch),
            "year_mismatch": len(year_mismatch),
            "unlinked_with_ps_in_db": len(unlinked_with_known_code),
            "unlinked_missing_ps_in_db": len(missing_target_ps),
            "order_number_collisions": len(colliding_orders),
        },
        "code_mismatch": code_mismatch[:200],
        "name_mismatch": name_mismatch[:200],
        "year_mismatch": year_mismatch[:200],
        "unlinked_with_ps_in_db": unlinked_with_known_code[:200],
        "unlinked_missing_ps_in_db_sample": missing_target_ps[:100],
        "order_number_collisions_sample": [
            {"order_token": token, "standards": items[:8]}
            for token, items in sorted(
                colliding_orders.items(), key=lambda x: -len(x[1])
            )[:40]
        ],
    }


def relink_and_audit(session: Session) -> dict:
    """Полная пересвязка + аудит. Используется админ-эндпоинтом и скриптом."""
    result = relink_all_qualifications(session, update_existing=True, commit=True)
    audit = audit_qualification_links(session)
    return {
        "status": "ok",
        "linked": result.linked,
        "skipped": result.skipped,
        "not_found": result.not_found,
        "by_method": result.by_method,
        "audit": audit,
    }
