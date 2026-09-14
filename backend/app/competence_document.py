"""Формирование DOCX-паспорта компетенции."""
from __future__ import annotations

import io
import re
from typing import Any

from .reference_data import (
    DESCRIPTOR_CATEGORIES,
    FORMATION_LEVEL_CODES,
    get_order_148n_indicators,
    normalize_descriptors,
    normalize_qualification_level_code,
)

KIND_LABELS = {
    "professional": "Профессиональная",
    "general": "Общекультурная",
    "universal": "Универсальная",
}

CATEGORY_LABELS = {
    "A": "Знания",
    "B": "Умения / интеллектуальные навыки",
    "C": "Практические навыки",
}

FORMATION_LABELS = {
    "базовый": "Базовый",
    "продвинутый": "Продвинутый",
    "экспертный": "Экспертный",
}


def _safe_filename(name: str) -> str:
    """ASCII-only имя для Content-Disposition (latin-1)."""
    text = (name or "competence").strip()
    # Транслит не обязателен: выкидываем всё не-ASCII
    ascii_only = text.encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", ascii_only)
    cleaned = re.sub(r"[^\w.\-]+", "_", cleaned, flags=re.ASCII)
    cleaned = re.sub(r"_+", "_", cleaned).strip("._")
    return (cleaned[:80] or "competence")


def competence_docx_filename(data: dict[str, Any]) -> str:
    cid = data.get("id")
    title = _safe_filename(str(data.get("name") or data.get("title") or "competence"))
    if cid:
        return f"competence_{cid}_{title}.docx" if title != "competence" else f"competence_{cid}.docx"
    return f"competence_{title}.docx"


def content_disposition_attachment(filename: str) -> str:
    """Безопасный Content-Disposition: ASCII filename + UTF-8 filename*."""
    from urllib.parse import quote

    ascii_name = _safe_filename(filename.replace(".docx", "")) + ".docx"
    # Если исходное имя уже ascii — одной формы достаточно
    if filename.isascii() and filename == ascii_name:
        return f'attachment; filename="{ascii_name}"'
    utf8_name = quote(filename if filename.endswith(".docx") else f"{filename}.docx")
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{utf8_name}"


def _structure_items(structure: Any, cat: str) -> list[str]:
    if not isinstance(structure, dict):
        return []
    raw = structure.get(cat) or []
    items: list[str] = []
    for item in raw:
        if isinstance(item, str) and item.strip():
            items.append(item.strip())
        elif isinstance(item, dict):
            text = str(item.get("text") or item.get("name") or "").strip()
            if text:
                items.append(text)
    return items


def _labor_functions(data: dict[str, Any]) -> list[dict[str, str]]:
    raw = data.get("labor_functions") or []
    result: list[dict[str, str]] = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        code = str(item.get("code") or "").strip()
        name = str(item.get("name") or "").strip()
        if code or name:
            result.append({"code": code, "name": name})
    return result


def _assessment_by_level(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    by_level: dict[str, dict[str, Any]] = {
        level: {"methods": [], "criteria": ""} for level in FORMATION_LEVEL_CODES
    }
    # New form shape: assessment_by_level
    abl = data.get("assessment_by_level")
    if isinstance(abl, dict):
        for level in FORMATION_LEVEL_CODES:
            block = abl.get(level) or {}
            methods = block.get("methods") or []
            by_level[level] = {
                "methods": [str(m).strip() for m in methods if str(m).strip()],
                "criteria": str(block.get("criteria") or "").strip(),
            }
        return by_level

    # Stored shape: assessment_tools = [{level, tool, criteria}, ...]
    tools = data.get("assessment_tools") or []
    for tool in tools:
        if not isinstance(tool, dict):
            continue
        level = str(tool.get("level") or "").strip()
        if level not in by_level:
            continue
        method = str(tool.get("tool") or tool.get("method") or "").strip()
        if method and method not in by_level[level]["methods"]:
            by_level[level]["methods"].append(method)
        criteria = str(tool.get("criteria") or "").strip()
        if criteria and not by_level[level]["criteria"]:
            by_level[level]["criteria"] = criteria
    return by_level


def build_competence_docx_bytes(data: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    for section in document.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    def set_run_font(run, size=11, bold=False):
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), "Times New Roman")

    def add_para(
        text: str = "",
        *,
        bold: bool = False,
        size: int = 11,
        center: bool = False,
        space_after: int = 6,
        space_before: int = 0,
    ):
        para = document.add_paragraph()
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(space_before)
        if text:
            run = para.add_run(text)
            set_run_font(run, size=size, bold=bold)
        return para

    def add_heading(text: str, level: int = 1):
        size = 14 if level == 1 else 12
        add_para(text, bold=True, size=size, space_before=12, space_after=8)

    def add_field(label: str, value: str):
        para = document.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        run_l = para.add_run(f"{label}: ")
        set_run_font(run_l, size=11, bold=True)
        run_v = para.add_run(value.strip() or "—")
        set_run_font(run_v, size=11, bold=False)

    def add_multiline(text: str):
        for line in (text or "—").splitlines() or ["—"]:
            add_para(line.strip() or "—", size=11, space_after=2)

    def set_cell_text(cell, text: str, *, bold: bool = False, size: int = 10, center: bool = False):
        cell.text = ""
        para = cell.paragraphs[0]
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.space_before = Pt(2)
        body = (text or "").strip() or "—"
        # Многострочный текст — отдельные абзацы в ячейке
        lines = body.splitlines() or ["—"]
        for i, line in enumerate(lines):
            target = para if i == 0 else cell.add_paragraph()
            if i > 0:
                target.paragraph_format.space_after = Pt(2)
                target.paragraph_format.space_before = Pt(0)
                if center:
                    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = target.add_run(line.strip() or "—")
            set_run_font(run, size=size, bold=bold)

    def shade_cell(cell, fill: str = "D9E2F3"):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    title = str(data.get("name") or data.get("title") or "Компетенция").strip()
    kind = str(data.get("competence_kind") or "professional")
    ql_raw = data.get("qualification_level") or data.get("qualification_level_code")
    ql_code = normalize_qualification_level_code(ql_raw)
    ql_label = (
        str(data.get("qualification_level_label") or "").strip()
        or (f"{ql_code}-й уровень" if ql_code else str(ql_raw or "").strip())
    )
    order_148n = str(data.get("order_148n_indicators") or "").strip()
    if not order_148n and ql_code:
        order_148n = get_order_148n_indicators(ql_code)

    descriptors = normalize_descriptors(data.get("descriptors"))
    structure = data.get("structure") or {}
    labor = _labor_functions(data)
    assessment = _assessment_by_level(data)

    # Title
    add_para("ПАСПОРТ КОМПЕТЕНЦИИ", bold=True, size=16, center=True, space_after=4)
    add_para(title, bold=True, size=14, center=True, space_after=12)
    if data.get("id"):
        add_para(f"ID: {data.get('id')}", size=10, center=True, space_after=12)

    # 1. General
    add_heading("1. Общая информация", 1)
    add_field("Название", title)
    add_field("Тип компетенции", KIND_LABELS.get(kind, kind))
    add_field("Отрасль", str(data.get("industry") or ""))
    add_field("Вид профессионального образования", str(data.get("education_kind") or ""))
    add_field("Уровень образования", str(data.get("education_level") or ""))
    add_field(
        "Профессия / должность (профобучение)",
        str(data.get("education_training_profession") or ""),
    )
    fgos_label = " ".join(
        part for part in (str(data.get("fgos_code") or "").strip(), str(data.get("fgos_name") or "").strip()) if part
    )
    add_field("ФГОС", fgos_label)
    add_field("Трудоёмкость", str(data.get("hours") or data.get("workload") or ""))
    add_field("Разработчик", str(data.get("developer") or ""))
    if data.get("status"):
        add_field("Статус", str(data.get("status")))
    add_para("Описание", bold=True, size=11, space_before=8, space_after=4)
    add_multiline(str(data.get("description") or ""))

    # 2. Standard & level
    add_heading("2. Профстандарт и уровень квалификации", 1)
    ps_name = str(data.get("prof_standard_name") or "").strip()
    ps_id = data.get("prof_standard_id")
    if ps_name:
        add_field("Профстандарт", ps_name)
    elif ps_id:
        add_field("Профстандарт ID", str(ps_id))
    else:
        add_field("Профстандарт", "Не указан" if kind == "professional" else "Не требуется")
    add_field("Уровень квалификации (приказ №148н)", ql_label)

    if order_148n:
        add_para("Показатели по приказу Минтруда №148н", bold=True, size=11, space_before=8, space_after=4)
        add_multiline(order_148n)

    if labor:
        add_para("Трудовые функции", bold=True, size=11, space_before=8, space_after=4)
        for item in labor:
            line = f"{item['code']} — {item['name']}".strip(" —")
            add_para(f"• {line}", size=11, space_after=2)
    else:
        add_field("Трудовые функции", "Не указаны")

    # 3. Structure
    add_heading("3. Структура A / B / C", 1)
    for cat in DESCRIPTOR_CATEGORIES:
        items = _structure_items(structure, cat)
        add_para(f"{cat}. {CATEGORY_LABELS.get(cat, cat)}", bold=True, size=11, space_before=6, space_after=4)
        if items:
            for idx, text in enumerate(items, 1):
                add_para(f"{idx}. {text}", size=11, space_after=2)
        else:
            add_para("Не заполнено", size=11, space_after=4)

    # 4. Descriptors — таблица: категории × уровни сформированности
    add_heading("4. Дескрипторы уровней сформированности", 1)
    desc_table = document.add_table(rows=1 + len(DESCRIPTOR_CATEGORIES), cols=1 + len(FORMATION_LEVEL_CODES))
    desc_table.style = "Table Grid"
    desc_table.autofit = True

    header_row = desc_table.rows[0]
    set_cell_text(header_row.cells[0], "Категория", bold=True, size=10, center=True)
    shade_cell(header_row.cells[0])
    for col_idx, level in enumerate(FORMATION_LEVEL_CODES, start=1):
        set_cell_text(
            header_row.cells[col_idx],
            FORMATION_LABELS.get(level, level),
            bold=True,
            size=10,
            center=True,
        )
        shade_cell(header_row.cells[col_idx])

    for row_idx, cat in enumerate(DESCRIPTOR_CATEGORIES, start=1):
        row = desc_table.rows[row_idx]
        set_cell_text(
            row.cells[0],
            f"{cat}. {CATEGORY_LABELS.get(cat, cat)}",
            bold=True,
            size=10,
        )
        shade_cell(row.cells[0], "F2F2F2")
        for col_idx, level in enumerate(FORMATION_LEVEL_CODES, start=1):
            text = (descriptors.get(cat) or {}).get(level) or ""
            set_cell_text(row.cells[col_idx], text, size=10)

    add_para("", space_after=6)

    # 5. Assessment
    add_heading("5. Оценочные средства", 1)
    for level in FORMATION_LEVEL_CODES:
        block = assessment[level]
        add_para(FORMATION_LABELS.get(level, level), bold=True, size=11, space_before=6, space_after=4)
        methods = block.get("methods") or []
        criteria = block.get("criteria") or ""
        if not methods and not criteria:
            add_para("Не заполнено", size=11, space_after=4)
            continue
        add_field("Методы", ", ".join(methods) if methods else "—")
        if criteria:
            add_para("Критерии и задания:", bold=True, size=10, space_after=2)
            add_multiline(criteria)

    # Optional resources / technologies
    techs = data.get("ed_technologies") or []
    if isinstance(techs, list) and techs:
        add_heading("6. Образовательные технологии", 1)
        add_para(", ".join(str(t) for t in techs if t), size=11)

    resources = data.get("resources") or []
    if isinstance(resources, list) and resources:
        add_heading("7. Ресурсы", 1)
        for item in resources:
            add_para(f"• {item}", size=11, space_after=2)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
