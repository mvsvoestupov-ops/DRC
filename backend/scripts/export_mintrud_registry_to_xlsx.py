"""
Экспорт реестра ПС с сайта Минтруда в Excel.

Колонки:
  Регистрационный номер
  Код профессионального стандарта
  Наименование
  Ответственная организация – разработчик
  Дата вступления в силу
  Дата утраты силы

Запуск:
  cd backend
  venv\\Scripts\\python.exe scripts\\export_mintrud_registry_to_xlsx.py
  venv\\Scripts\\python.exe scripts\\export_mintrud_registry_to_xlsx.py --page-size 20
  venv\\Scripts\\python.exe scripts\\export_mintrud_registry_to_xlsx.py -o "C:\\path\\reestr.xlsx"
"""
from __future__ import annotations

import argparse
import os
import sys
import time
import warnings
from datetime import date

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)

from app.parser import REGISTRY_BASE_URL, fetch_registry_page_items

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")

HEADERS = [
    "Регистрационный номер",
    "Код профессионального стандарта",
    "Наименование",
    "Ответственная организация – разработчик",
    "Дата вступления в силу",
    "Дата утраты силы",
]


def fetch_all_registry_rows(page_size: int, max_pages: int | None = None) -> tuple[list[dict], int | None]:
    all_items: dict[str, dict] = {}
    page_total: int | None = None
    page = 1
    empty_streak = 0

    print(f"Загрузка реестра с сайта Минтруда (по {page_size} на страницу)...")
    print(f"URL: {REGISTRY_BASE_URL}")
    while empty_streak < 2:
        if max_pages is not None and page > max_pages:
            break
        print(f"  страница {page}...", end=" ", flush=True)
        items, total = fetch_registry_page_items(page, page_size)
        if total and page_total is None:
            page_total = total
            print(f"всего на сайте: {total}", end=" ")
        if not items:
            print("пусто")
            empty_streak += 1
            page += 1
            continue

        empty_streak = 0
        new = 0
        for item in items:
            key = item.get("element_id") or f"{item.get('reg_number')}|{item.get('ps_code')}"
            if key not in all_items:
                all_items[key] = item
                new += 1
        print(f"строк={len(items)}, новых={new}, уникальных={len(all_items)}")
        page += 1
        time.sleep(0.4)

        if page_total and len(all_items) >= page_total:
            break

    rows = list(all_items.values())
    rows.sort(
        key=lambda item: (
            int(item["reg_number"]) if str(item.get("reg_number", "")).isdigit() else 99999,
            item.get("ps_code") or "",
        )
    )
    return rows, page_total


def default_output_path() -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    return os.path.join(OUTPUT_DIR, f"reestr_mintrud_{date.today().isoformat()}.xlsx")


def write_xlsx(rows: list[dict], path: str) -> None:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Реестр ПС"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for item in rows:
        ws.append(
            [
                item.get("reg_number") or "",
                item.get("ps_code") or "",
                item.get("name") or "",
                item.get("developer") or "",
                item.get("effective_date") or "",
                item.get("expiration_date") or "",
            ]
        )

    widths = [18, 22, 60, 50, 20, 20]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    wb.save(path)
    wb.close()


def main() -> int:
    warnings.filterwarnings("ignore", message="Unverified HTTPS request")

    parser = argparse.ArgumentParser(description="Экспорт реестра ПС с сайта Минтруда в Excel")
    parser.add_argument(
        "--page-size",
        type=int,
        default=20,
        help="Записей на страницу (на сайте по умолчанию 20)",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Путь к выходному XLSX (по умолчанию scripts/output/reestr_mintrud_YYYY-MM-DD.xlsx)",
    )
    parser.add_argument(
        "--max-pages",
        type=int,
        help="Ограничить число страниц (для проверки, например --max-pages 1)",
    )
    args = parser.parse_args()

    rows, site_total = fetch_all_registry_rows(args.page_size, max_pages=args.max_pages)
    if not rows:
        print("Не удалось получить записи с сайта.")
        return 1

    output_path = os.path.abspath(args.output or default_output_path())
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    write_xlsx(rows, output_path)

    print("\n" + "=" * 60)
    print("ЭКСПОРТ ЗАВЕРШЁН")
    print("=" * 60)
    print(f"Записей: {len(rows)}", end="")
    if site_total:
        print(f" (на сайте указано: {site_total})")
    else:
        print()
    print(f"Файл: {output_path}")

    sample = rows[0]
    print("\nПример первой строки:")
    print(f"  reg={sample.get('reg_number')} code={sample.get('ps_code')}")
    print(f"  name={(sample.get('name') or '')[:70]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
