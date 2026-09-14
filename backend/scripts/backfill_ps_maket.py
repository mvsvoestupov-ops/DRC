"""
Backfill макета ПС: метаданные реестра + XML Минтруда + classinform HTML.

Примеры:
  venv\\Scripts\\python.exe scripts\\backfill_ps_maket.py --limit 20
  venv\\Scripts\\python.exe scripts\\backfill_ps_maket.py --registry-only
  venv\\Scripts\\python.exe scripts\\backfill_ps_maket.py --only-gaps
  venv\\Scripts\\python.exe scripts\\backfill_ps_maket.py --reg 1582
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
import warnings
from collections import Counter

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND_DIR)

from app.db import SessionLocal
from app.db.raw_models import StandardRaw
from app.db.schema import ensure_maket_columns, ensure_ps_status_columns
from app.db_operations import apply_registry_metadata, save_raw_standard
from app.classinform_parser import build_ps_index, load_ps_from_classinform
from app.parser import parse_xml, download_bulk_xml_chunk

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")
REPORT_JSON = os.path.join(OUTPUT_DIR, "ps_maket_backfill_report.json")
REPORT_TXT = os.path.join(OUTPUT_DIR, "ps_maket_backfill_report.txt")
DOWNLOADS = os.path.join(BACKEND_DIR, "downloads")


def _has_education(std: StandardRaw) -> bool:
    return any(getattr(gf, "education_training", None) for gf in std.generalized_functions)


def _gap_flags(std: StandardRaw) -> list[str]:
    gaps = []
    if not getattr(std, "ps_code", None):
        gaps.append("ps_code")
    if not getattr(std, "developer_org", None):
        gaps.append("developer_org")
    if not _has_education(std):
        gaps.append("education_training")
    if not getattr(std, "source_html", None) and not getattr(std, "source_xml", None):
        gaps.append("source")
    if not getattr(std, "abbreviations", None):
        gaps.append("abbreviations")
    return gaps


def _infer_ps_code(std: StandardRaw) -> str | None:
    if getattr(std, "ps_code", None):
        return std.ps_code
    eid = std.element_id or ""
    if eid.startswith("classinform:"):
        return eid.split(":", 1)[1]
    # иногда professional_area_code + хвост из имени ненадёжны — только явный XX.XXX
    return None


def backfill_registry_from_xlsx(session, xlsx_path: str) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"updated": 0}
    header = [str(c or "").strip().lower() for c in rows[0]]

    def find_col(*needles):
        for i, h in enumerate(header):
            for n in needles:
                if n in h:
                    return i
        return None

    i_reg = find_col("регистрац", "рег")
    i_code = find_col("код")
    i_dev = find_col("разработ")
    i_eff = find_col("дата введения", "введен")
    i_exp = find_col("дата утрат", "утратил")
    updated = 0
    for row in rows[1:]:
        if not row or i_reg is None:
            continue
        reg = str(row[i_reg] or "").strip()
        if not reg:
            continue
        ps_code = str(row[i_code] or "").strip() if i_code is not None else None
        developer = str(row[i_dev] or "").strip() if i_dev is not None else None
        effective = str(row[i_eff] or "").strip() if i_eff is not None else None
        expiration = str(row[i_exp] or "").strip() if i_exp is not None else None
        if apply_registry_metadata(
            session,
            reg,
            ps_code=ps_code or None,
            developer_org=developer or None,
            effective_date=effective or None,
            expiration_date=expiration or None,
        ):
            updated += 1
    return {"updated": updated, "source": xlsx_path}


def try_reload_xml(std: StandardRaw) -> tuple[bool, str]:
    eid = (std.element_id or "").strip()
    if not eid.isdigit():
        return False, "no_numeric_element_id"
    os.makedirs(DOWNLOADS, exist_ok=True)
    path = os.path.join(DOWNLOADS, f"backfill_{eid}.xml")
    try:
        download_bulk_xml_chunk([eid], path)
        content = open(path, "rb").read()
        if b"403" in content[:200] or len(content) < 200:
            return False, "xml_too_small_or_forbidden"
        standard = parse_xml(content, element_id=eid)
        if std.reg_number and standard.registration_number != std.reg_number:
            standard.registration_number = std.reg_number
        session = SessionLocal()
        try:
            save_raw_standard(session, standard, element_id=eid)
        finally:
            session.close()
        return True, "xml_ok"
    except Exception as exc:
        return False, f"xml_error:{exc}"


def try_reload_classinform(std: StandardRaw, index: dict, force: bool = False) -> tuple[bool, str]:
    ps_code = _infer_ps_code(std)
    if not ps_code:
        # попробуем из индекса по имени — слишком дорого; требуем код
        area = std.professional_area_code
        return False, "no_ps_code"
    try:
        ok = load_ps_from_classinform(ps_code, reg_number=std.reg_number, index=index, force=True)
        return (True, "classinform_ok") if ok else (False, "classinform_fail")
    except Exception as exc:
        return False, f"classinform_error:{exc}"


def maket_completeness(std: StandardRaw) -> dict:
    gaps = _gap_flags(std)
    p0 = ["ps_code", "developer_org", "education_training"]
    p0_ok = all(g not in gaps for g in p0)
    return {"reg_number": std.reg_number, "gaps": gaps, "p0_ok": p0_ok}


def main() -> int:
    warnings.filterwarnings("ignore", message="Unverified HTTPS request")
    parser = argparse.ArgumentParser(description="Backfill макета ПС")
    parser.add_argument("--limit", type=int, default=0, help="Ограничить число ПС (0 = все)")
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--reg", type=str, default="", help="Только один рег. номер")
    parser.add_argument("--registry-only", action="store_true")
    parser.add_argument("--only-gaps", action="store_true", help="Только ПС с пробелами P0")
    parser.add_argument("--skip-xml", action="store_true")
    parser.add_argument("--skip-classinform", action="store_true")
    parser.add_argument("--delay", type=float, default=0.4)
    parser.add_argument(
        "--xlsx",
        default=os.path.join(OUTPUT_DIR, "reestr_mintrud_2026-08-31.xlsx"),
        help="XLSX реестра для developer/дат/ps_code",
    )
    args = parser.parse_args()

    for col in ensure_ps_status_columns() + ensure_maket_columns():
        print(f"DB migration: {col}")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    session = SessionLocal()
    stats = Counter()
    details = []

    try:
        if os.path.exists(args.xlsx):
            print(f"Реестр XLSX: {args.xlsx}")
            reg_stats = backfill_registry_from_xlsx(session, args.xlsx)
            stats["registry_updated"] = reg_stats.get("updated", 0)
            print(f"  обновлено метаданных: {stats['registry_updated']}")
        else:
            print(f"XLSX не найден ({args.xlsx}) — пропускаем registry metadata")

        if args.registry_only:
            _write_report(session, stats, details)
            return 0

        q = session.query(StandardRaw).order_by(StandardRaw.reg_number)
        if args.reg:
            q = q.filter(StandardRaw.reg_number == args.reg)
        standards = q.offset(args.offset).all()
        if args.limit:
            standards = standards[: args.limit]

        index = None
        if not args.skip_classinform:
            print("Индекс classinform...")
            index = build_ps_index()

        for i, std in enumerate(standards, start=1):
            gaps = _gap_flags(std)
            if args.only_gaps and not any(g in gaps for g in ("education_training", "developer_org", "source", "ps_code")):
                stats["skipped_complete"] += 1
                continue

            print(f"[{i}/{len(standards)}] reg={std.reg_number} gaps={gaps}")
            entry = {"reg_number": std.reg_number, "before_gaps": gaps, "actions": []}

            # XML
            if not args.skip_xml and (std.element_id or "").isdigit():
                ok, msg = try_reload_xml(std)
                entry["actions"].append(msg)
                stats[msg] += 1
                time.sleep(args.delay)
                session.expire_all()
                std = session.query(StandardRaw).filter(StandardRaw.reg_number == entry["reg_number"]).first()

            # classinform если дыры в III/IV или нет source
            need_ci = False
            if std:
                g2 = _gap_flags(std)
                need_ci = any(g in g2 for g in ("education_training", "abbreviations", "source")) or (
                    std.element_id or ""
                ).startswith("classinform:")
            if not args.skip_classinform and need_ci and index is not None:
                ok, msg = try_reload_classinform(std, index, force=True)
                entry["actions"].append(msg)
                stats[msg] += 1
                time.sleep(args.delay)
                session.expire_all()
                std = session.query(StandardRaw).filter(StandardRaw.reg_number == entry["reg_number"]).first()

            if std:
                after = maket_completeness(std)
                entry["after"] = after
                stats["p0_ok" if after["p0_ok"] else "p0_gap"] += 1
            details.append(entry)

        _write_report(session, stats, details)
        print("Stats:", dict(stats))
        print(f"Report: {REPORT_TXT}")
        return 0
    finally:
        session.close()


def _write_report(session, stats: Counter, details: list) -> None:
    all_std = session.query(StandardRaw).all()
    completeness = [maket_completeness(s) for s in all_std]
    p0_ok = sum(1 for c in completeness if c["p0_ok"])
    gap_counter = Counter()
    for c in completeness:
        gap_counter.update(c["gaps"])

    report = {
        "stats": dict(stats),
        "totals": {
            "standards": len(all_std),
            "p0_ok": p0_ok,
            "p0_ok_pct": round(100.0 * p0_ok / len(all_std), 2) if all_std else 0,
            "gap_counts": dict(gap_counter),
        },
        "details": details[-500:],
    }
    with open(REPORT_JSON, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

    lines = [
        "ОТЧЁТ BACKFILL МАКЕТА ПС",
        "=" * 50,
        f"Всего ПС: {len(all_std)}",
        f"P0 заполнены (ps_code+developer+education): {p0_ok} ({report['totals']['p0_ok_pct']}%)",
        f"Пробелы: {dict(gap_counter)}",
        f"Действия прогона: {dict(stats)}",
    ]
    with open(REPORT_TXT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


if __name__ == "__main__":
    raise SystemExit(main())
